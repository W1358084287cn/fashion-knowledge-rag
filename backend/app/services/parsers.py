# -*- coding: utf-8 -*-
"""
多源文件解析器
------------------------------------------------------------
支持：TXT / Markdown / Word / Excel / PPT / 普通 PDF / 扫描 PDF（PaddleOCR）/ 客服聊天记录
统一输出纯文本，后续交给 Agent1 清洗、切片、打标签。

说明：
- PaddleOCR 首次使用会自动下载中文模型，且体积较大，因此采用“懒加载”，
  只有真正遇到扫描件时才初始化，不影响普通文件的处理速度。
- 若环境未安装 PaddleOCR，会给出明确的中文错误提示，而不是静默失败。
"""
import csv
import json
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

from ..config import settings


class ParseError(Exception):
    """文件解析异常：由上层捕获后写入 kb_files.error_msg，展示给管理员。"""


@dataclass
class ParseResult:
    """解析结果。"""
    text: str
    source_type: str = "文档"          # 文档 / 聊天记录 / 表格数据
    meta: dict = field(default_factory=dict)


# ------------------------------------------------------------------
# 通用工具
# ------------------------------------------------------------------
def _read_text(path: str) -> str:
    """按常见中文编码依次尝试读取文本文件。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ParseError("文件编码无法识别，请另存为 UTF-8 后重试")


def _join_rows(headers: List[str], rows: List[List[str]]) -> str:
    """把表格行转换为自然语言句子，便于切片与语义检索。"""
    lines = []
    for row in rows:
        parts = []
        for i, value in enumerate(row):
            value = str(value).strip()
            if not value or value.lower() == "nan":
                continue
            header = headers[i] if i < len(headers) and headers[i] else f"字段{i + 1}"
            parts.append(f"{header}：{value}")
        if parts:
            lines.append("；".join(parts) + "。")
    return "\n".join(lines)


# ------------------------------------------------------------------
# Word / PPT
# ------------------------------------------------------------------
def parse_docx(path: str) -> ParseResult:
    from docx import Document

    doc = Document(path)
    blocks: List[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            blocks.append(para.text.strip())
    # 表格单独展开，避免信息丢失
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            blocks.append(_join_rows(rows[0], rows[1:]))
    return ParseResult(text="\n".join(blocks), source_type="文档")


def parse_pptx(path: str) -> ParseResult:
    from pptx import Presentation

    prs = Presentation(path)
    blocks: List[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_texts: List[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        slide_texts.append(t)
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                if rows:
                    slide_texts.append(_join_rows(rows[0], rows[1:]))
        if slide_texts:
            blocks.append(f"第{idx}页：\n" + "\n".join(slide_texts))
    return ParseResult(text="\n".join(blocks), source_type="文档")


# ------------------------------------------------------------------
# Excel / CSV
# ------------------------------------------------------------------
def parse_excel(path: str) -> ParseResult:
    import pandas as pd

    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Excel 解析失败：{exc}") from exc

    blocks: List[str] = []
    for sheet_name, df in sheets.items():
        df = df.fillna("")
        if df.empty:
            continue
        headers = [str(c).strip() for c in df.columns]
        rows = df.astype(str).values.tolist()
        blocks.append(f"工作表【{sheet_name}】：\n" + _join_rows(headers, rows))
    return ParseResult(text="\n".join(blocks), source_type="表格数据")


def parse_csv(path: str) -> ParseResult:
    import pandas as pd

    raw = _read_text(path)
    try:
        df = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    except Exception:
        # 兼容 GBK 编码的客服导出表
        try:
            df = pd.read_csv(path, dtype=str, encoding="gb18030")
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"CSV 解析失败：{exc}") from exc
    df = df.fillna("")
    headers = [str(c).strip() for c in df.columns]

    # 识别客服聊天记录表：包含发送人 / 角色 / 内容等列
    lower_headers = [h.lower() for h in headers]
    speaker_keys = ["发送人", "发言人", "角色", "sender", "role", "客服", "用户"]
    content_keys = ["内容", "消息", "content", "message", "text"]
    is_chat = (
        any(any(k in h for k in speaker_keys) for h in lower_headers)
        and any(any(k in h for k in content_keys) for h in lower_headers)
    )
    if is_chat:
        lines = []
        for _, row in df.iterrows():
            values = [str(v).strip() for v in row.tolist()]
            if any(values):
                lines.append(" | ".join(v for v in values if v))
        return ParseResult(text="\n".join(lines), source_type="聊天记录", meta={"raw_length": len(raw)})

    rows = df.astype(str).values.tolist()
    return ParseResult(text=_join_rows(headers, rows), source_type="表格数据")


# ------------------------------------------------------------------
# 客服聊天记录（TXT / JSON）
# ------------------------------------------------------------------
def parse_chat_txt(path: str) -> ParseResult:
    """
    客服聊天记录 TXT：
    常见格式 `[2026-01-01 10:00:00] 用户：这件大衣起球吗？`
    清洗时去掉时间戳，仅保留“说话人：内容”，降低噪声。
    """
    text = _read_text(path)
    cleaned_lines: List[str] = []
    timestamp_pattern = re.compile(
        r"^\s*[\[\(]?\d{4}[-/年]\d{1,2}[-/月]\d{1,2}.*?[\]\)]?\s*"
    )
    for line in text.splitlines():
        line = timestamp_pattern.sub("", line).strip()
        if line:
            cleaned_lines.append(line)
    return ParseResult(text="\n".join(cleaned_lines), source_type="聊天记录")


def parse_chat_json(path: str) -> ParseResult:
    """客服聊天记录 JSON：兼容 [{role, content}] 或 {messages:[...]}。"""
    try:
        data = json.loads(_read_text(path))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"聊天记录 JSON 解析失败：{exc}") from exc

    if isinstance(data, dict):
        data = data.get("messages") or data.get("data") or []
    if not isinstance(data, list):
        raise ParseError("聊天记录 JSON 结构不支持，请使用数组格式")

    lines = []
    for item in data:
        if isinstance(item, dict):
            speaker = item.get("role") or item.get("sender") or item.get("name") or "用户"
            content = item.get("content") or item.get("message") or item.get("text") or ""
            if str(content).strip():
                lines.append(f"{speaker}：{str(content).strip()}")
        elif isinstance(item, str) and item.strip():
            lines.append(item.strip())
    return ParseResult(text="\n".join(lines), source_type="聊天记录")


# ------------------------------------------------------------------
# PDF（普通 + 扫描件 OCR）
# ------------------------------------------------------------------
_ocr_engine = None


def _get_ocr_engine():
    """懒加载 PaddleOCR：只在扫描件场景初始化，避免拖慢整体启动。"""
    global _ocr_engine
    if _ocr_engine is not None:
        return _ocr_engine
    try:
        from paddleocr import PaddleOCR
    except Exception as exc:  # noqa: BLE001
        raise ParseError(
            "当前环境未安装 PaddleOCR，无法识别扫描版 PDF。"
            "请执行 pip install -r requirements-ocr.txt 后重启服务。"
        ) from exc
    try:
        # show_log=False 关闭 paddle 的冗余日志；lang='ch' 使用中文模型
        _ocr_engine = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PaddleOCR 初始化失败：{exc}") from exc
    return _ocr_engine


def _ocr_pdf_page(page) -> str:
    """将 PDF 单页渲染为图片后交给 PaddleOCR 识别。"""
    import fitz  # PyMuPDF
    import numpy as np

    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if pix.n >= 3:
        img = img[:, :, :3][:, :, ::-1]  # RGB -> BGR，PaddleOCR 默认按 BGR 处理

    engine = _get_ocr_engine()
    result = engine.ocr(img, cls=True) or []
    lines: List[str] = []
    for page_result in result:
        for line in page_result or []:
            try:
                lines.append(line[1][0])
            except (IndexError, TypeError):
                continue
    return "\n".join(lines)


def parse_pdf(path: str) -> ParseResult:
    """
    PDF 解析策略：
    1. 先用 pdfplumber 抽取文本层（速度快、准确率高）
    2. 若整页几乎没有文字（扫描件），用 PyMuPDF 渲染 + PaddleOCR 兜底
    """
    import pdfplumber

    text_pages: List[str] = []
    scanned_pages: List[int] = []

    try:
        with pdfplumber.open(path) as pdf:
            total_pages = len(pdf.pages)
            for idx, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                text_pages.append(page_text)
                if len(page_text.strip()) < settings.OCR_MIN_CHARS_PER_PAGE:
                    scanned_pages.append(idx)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 解析失败：{exc}") from exc

    # 扫描页 OCR（最多处理 30 页，防止超大扫描件把演示环境拖死）
    if scanned_pages:
        try:
            import fitz

            max_pages = min(len(scanned_pages), 30)
            with fitz.open(path) as doc:
                for i in scanned_pages[:max_pages]:
                    if i < doc.page_count:
                        ocr_text = _ocr_pdf_page(doc[i])
                        if ocr_text.strip():
                            text_pages[i] = ocr_text
        except ParseError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"扫描件 OCR 失败：{exc}") from exc

    text = "\n\n".join(t for t in text_pages if t and t.strip())
    if not text.strip():
        raise ParseError("PDF 中未提取到有效文字，请确认文件内容或安装 OCR 依赖")
    return ParseResult(text=text, source_type="文档", meta={"pages": total_pages})


# ------------------------------------------------------------------
# SQL 数据文件
# ------------------------------------------------------------------
_SQL_COMMENT_BLOCK = re.compile(r"/\*[\s\S]*?\*/")
_SQL_COMMENT_LINE = re.compile(r"--[^\n]*")
_CREATE_TABLE = re.compile(
    r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?[`\"]?([\w.]+)[`\"]?\s*\(([\s\S]*?)\)\s*;",
    re.IGNORECASE,
)
_INSERT = re.compile(
    r"INSERT\s+INTO\s+[`\"]?([\w.]+)[`\"]?\s*(?:\(([^)]*)\))?\s*VALUES\s*([\s\S]*?);",
    re.IGNORECASE,
)


def _split_sql_tuples(values: str) -> List[str]:
    """按最外层括号拆分 VALUES 后的 (..),(..) 组，正确处理引号内的逗号。"""
    tuples: List[str] = []
    depth = 0
    quote = ""
    buffer = ""
    for ch in values:
        if quote:
            buffer += ch
            if ch == quote:
                quote = ""
            continue
        if ch in ("'", '"'):
            quote = ch
            buffer += ch
            continue
        if ch == "(":
            depth += 1
            if depth == 1:
                buffer = ""
            else:
                buffer += ch
            continue
        if ch == ")":
            depth -= 1
            if depth == 0:
                tuples.append(buffer)
                buffer = ""
            else:
                buffer += ch
            continue
        if depth >= 1:
            buffer += ch
    return tuples


def _split_sql_fields(raw: str) -> List[str]:
    """拆分一行 SQL 字段，忽略引号内的逗号。"""
    fields: List[str] = []
    quote = ""
    buffer = ""
    for ch in raw:
        if quote:
            buffer += ch
            if ch == quote:
                quote = ""
            continue
        if ch in ("'", '"'):
            quote = ch
            buffer += ch
            continue
        if ch == ",":
            fields.append(buffer.strip())
            buffer = ""
        else:
            buffer += ch
    if buffer.strip():
        fields.append(buffer.strip())
    return fields


def _unquote_sql(value: str) -> str:
    """去掉 SQL 字符串引号并还原常见转义。"""
    value = value.strip()
    if len(value) >= 2 and value[0] in ("'", '"') and value[-1] == value[0]:
        value = value[1:-1]
    value = value.replace("\\'", "'").replace('\\"', '"').replace("\\n", " ").replace("\\\\", "\\")
    if value.upper() == "NULL":
        return ""
    return value.strip()


def parse_sql(path: str) -> ParseResult:
    """
    SQL 文件解析：
    - 去掉 -- 与 /* */ 注释
    - CREATE TABLE：展开表名与字段说明
    - INSERT INTO：把每行数据转成“字段：值；...”的自然语言句子
    这样数据库导出文件也能直接进入向量库与知识图谱。
    """
    raw = _read_text(path)
    text = _SQL_COMMENT_BLOCK.sub(" ", raw)
    text = _SQL_COMMENT_LINE.sub(" ", text)

    blocks: List[str] = []

    # 1. 建表语句 -> 字段说明
    for match in _CREATE_TABLE.finditer(text):
        table = match.group(1).strip("`\"")
        body = match.group(2)
        columns = []
        for line in body.splitlines():
            line = line.strip().strip(",")
            if not line:
                continue
            upper = line.upper()
            if upper.startswith(("PRIMARY ", "UNIQUE ", "KEY ", "INDEX ", "CONSTRAINT ", "FOREIGN ")):
                continue
            col = re.split(r"\s+", line, maxsplit=1)[0].strip("`\"")
            if col:
                columns.append(col)
        if columns:
            blocks.append(f"数据表【{table}】包含字段：{'、'.join(columns)}。")

    # 2. 插入语句 -> 逐行数据
    for match in _INSERT.finditer(text):
        table = match.group(1).strip("`\"")
        columns = [
            c.strip().strip("`\"")
            for c in _split_sql_fields(match.group(2) or "")
            if c.strip()
        ]
        for row in _split_sql_tuples(match.group(3)):
            values = [_unquote_sql(v) for v in _split_sql_fields(row)]
            pairs = []
            for idx, value in enumerate(values):
                if not value:
                    continue
                header = columns[idx] if idx < len(columns) and columns[idx] else f"字段{idx + 1}"
                pairs.append(f"{header}：{value}")
            if pairs:
                blocks.append(f"数据表【{table}】：" + "；".join(pairs) + "。")

    result_text = "\n".join(blocks).strip()
    if not result_text:
        # 没有可识别的建表 / 插入语句时，回退为原文，交给清洗模块继续处理
        result_text = text.strip()
    if not result_text:
        raise ParseError("SQL 文件中未解析到有效数据")
    return ParseResult(text=result_text, source_type="表格数据",
                       meta={"tables": len(blocks)})


# ------------------------------------------------------------------
# 统一入口
# ------------------------------------------------------------------
def parse_file(path: str, ext: str, file_name: str = "") -> ParseResult:
    """
    根据扩展名分发到具体解析器。
    ext 统一小写且带点，如 `.pdf`。
    """
    ext = ext.lower()
    name_lower = (file_name or os.path.basename(path)).lower()
    is_chat_name = any(k in name_lower for k in ("聊天", "chat", "会话", "客服记录"))

    if ext in (".txt", ".md"):
        result = parse_chat_txt(path) if is_chat_name else ParseResult(text=_read_text(path))
    elif ext == ".docx":
        result = parse_docx(path)
    elif ext == ".pptx":
        result = parse_pptx(path)
    elif ext in (".xlsx", ".xls"):
        result = parse_excel(path)
    elif ext == ".csv":
        result = parse_csv(path)
    elif ext == ".json":
        result = parse_chat_json(path)
    elif ext == ".pdf":
        result = parse_pdf(path)
    elif ext == ".sql":
        result = parse_sql(path)
    else:
        raise ParseError(f"暂不支持的文件类型：{ext}")

    # 文件名中包含“聊天”关键字时，强制标记为聊天记录来源
    if is_chat_name and result.source_type == "文档":
        result.source_type = "聊天记录"

    if not result.text.strip():
        raise ParseError("文件内容为空或未解析到有效文本")
    return result
